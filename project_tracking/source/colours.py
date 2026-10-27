"""One clear colour per status and priority, shared by the three workbooks.

Only the Status / Priority / Result cells are coloured (not whole rows), with bold text, so a client or
project manager can scan a column. Conditional formatting: when someone changes a value in Excel, the
colour follows. Each palette maps value -> (fill colour, text colour).
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

WHITE, INK = "FFFFFF", "1F2A27"
GREEN, TEAL, BLUE, ORANGE, YELLOW, PURPLE, GREY, LIGHT_GREY, RED = (
	"2E7D32", "00897B", "1565C0", "EF6C00", "F9A825", "6A1B9A", "757575", "B0BEC5", "C62828")

TASK_STATUS = {
	"Done": (GREEN, WHITE),
	"Done – client UAT pending": (TEAL, WHITE),
	"In progress": (BLUE, WHITE),
	"Pending – decision needed": (ORANGE, WHITE),
	"Pending – data / setup": (YELLOW, INK),
	"Pending – development": (PURPLE, WHITE),
	"Pending – other team": (GREY, WHITE),
	"Planned": (LIGHT_GREY, INK),
}
BUG_STATUS = {
	"Fixed": (GREEN, WHITE),
	"Open": (RED, WHITE),
	"Open – other team": (GREY, WHITE),
	"Open – check server": (YELLOW, INK),
}
CHECKLIST_STATUS = {
	"Open": (ORANGE, WHITE),
	"In progress": (BLUE, WHITE),
	"Done": (GREEN, WHITE),
	"Blocked": (RED, WHITE),
	"Not needed": (LIGHT_GREY, INK),
}
TEST_RESULT = {
	"Pass": (GREEN, WHITE),
	"Fail": (RED, WHITE),
	"Blocked": (YELLOW, INK),
	"Not run": (LIGHT_GREY, INK),
	"N/A": (LIGHT_GREY, INK),
}
TESTER_BUG_STATUS = {
	"Open": (RED, WHITE),
	"Fixed": (TEAL, WHITE),
	"Retest": (YELLOW, INK),
	"Closed": (GREEN, WHITE),
	"Not a bug": (LIGHT_GREY, INK),
}
PRIORITY = {
	"Critical": ("B71C1C", WHITE),
	"High": ("E64A19", WHITE),
	"Medium": ("FBC02D", INK),
	"Low": ("90A4AE", INK),
}


def colour_column(ws, col, first, last, palette):
	"""Colour the cells of column `col` (rows first..last) by their value."""
	for value, (fill, text) in palette.items():
		ws.conditional_formatting.add(f"{col}{first}:{col}{last}", FormulaRule(
			formula=[f'${col}{first}="{value}"'], stopIfTrue=True,
			# Excel draws a conditional solid fill from bgColor; set both so every viewer shows it.
			fill=PatternFill("solid", fgColor=fill, bgColor=fill),
			font=Font(color=text, bold=True)))


def badge(cell, palette, value, font_name="Arial"):
	"""The same colours on a fixed cell (legends and summary labels)."""
	fill, text = palette[value]
	cell.fill = PatternFill("solid", fgColor=fill)
	cell.font = Font(name=font_name, size=10, bold=True, color=text)
