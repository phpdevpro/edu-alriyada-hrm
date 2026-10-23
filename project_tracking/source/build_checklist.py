"""Build Medical_HRMS_Jameah_Checklist.xlsx from checklist_data.py (run build_all.py)."""
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from changelog import CHANGELOG
from checklist_data import CHECKLIST, STATUSES
from colours import CHECKLIST_STATUS, colour_column

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "Medical_HRMS_Jameah_Checklist.xlsx")
AS_OF = CHANGELOG[0][0]
F = "Arial"
INK, MUTED, ACCENT = "1F2A27", "5D6B66", "0D6B4F"
thin = Side(style="thin", color="D5DDDA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=ACCENT)
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")


def header(ws, names, widths):
	for i, (h, w) in enumerate(zip(names, widths), 1):
		c = ws.cell(row=1, column=i, value=h)
		c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
		c.fill = HEAD_FILL
		c.alignment = Alignment(vertical="center", wrap_text=True)
		c.border = BORDER
		ws.column_dimensions[get_column_letter(i)].width = w
	ws.row_dimensions[1].height = 30
	ws.freeze_panes = "C2"


def cell(ws, r, c, v, wrap=False, bold=False, fill=None):
	x = ws.cell(row=r, column=c, value=v)
	x.font = Font(name=F, size=10, color=INK, bold=bold)
	x.alignment = Alignment(vertical="top", wrap_text=wrap)
	x.border = BORDER
	if fill:
		x.fill = fill
	return x


wb = Workbook()

# ---------------------------------------------------------------- Read Me
ws = wb.active
ws.title = "Read Me"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 110
lines = [
	("t", "Medical HRMS & Jameah — Checklist"),
	("s", f"Al-Riyada College for Health Sciences · last updated {AS_OF}"),
	("", ""),
	("p", "What still has to be done or checked outside the code: deploying to the live server, checks only the live server can answer, decisions only the college or the Ministry can make, HR data entry, and reviews."),
	("p", "Work through 'Checklist'. When an item is done, set Status to Done and write the date in 'Done on'. Each row shows the task or decision it belongs to (Ref) in the Task Tracker."),
	("p", "Every change to an item updates its 'Last updated' date and gets a line in 'Change Log' (developers: project_tracking/source, then build_all.py)."),
	("p", "Companion files in the same folder: Medical_HRMS_Jameah_Task_Tracker.xlsx (all tasks, bugs and decisions) and Medical_HRMS_Jameah_Test_Cases.xlsx (for the testing team)."),
]
for r, (kind, v) in enumerate(lines, 1):
	c = ws.cell(row=r, column=2, value=v)
	if kind == "t":
		c.font = Font(name=F, bold=True, size=18, color=ACCENT)
	elif kind == "s":
		c.font = Font(name=F, size=11, color=MUTED)
	else:
		c.font = Font(name=F, size=10, color=INK)
		c.alignment = Alignment(wrap_text=True, vertical="top")
		ws.row_dimensions[r].height = 30

# ---------------------------------------------------------------- Checklist
wc = wb.create_sheet("Checklist")
cols = ["ID", "Group", "Item", "How / where", "Owner", "Needed by", "Ref (tracker)", "Status", "Done on", "Added on", "Last updated", "Notes"]
header(wc, cols, [8, 26, 34, 70, 22, 16, 14, 12, 12, 12, 12, 30])
for i, row in enumerate(CHECKLIST, 2):
	for j, v in enumerate(list(row) + [""], 1):
		cell(wc, i, j, v, wrap=j in (3, 4, 12), bold=j == 1, fill=INPUT_FILL if j in (8, 9, 12) else None)
last = len(CHECKLIST) + 1
wc.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
wc.add_data_validation(dv)
dv.add(f"H2:H{last + 100}")
colour_column(wc, "H", 2, last + 100, CHECKLIST_STATUS)

# ---------------------------------------------------------------- Summary
wsum = wb.create_sheet("Summary", 1)
wsum.sheet_view.showGridLines = False
wsum["A1"] = "Summary"
wsum["A1"].font = Font(name=F, bold=True, size=16, color=ACCENT)
wsum["A2"] = "Counts are formulas over the Checklist sheet; they update as Status changes."
wsum["A2"].font = Font(name=F, size=10, color=MUTED)
groups = list(dict.fromkeys(r[1] for r in CHECKLIST))
head = ["Group", "Items"] + STATUSES
for i, (h, w) in enumerate(zip(head, [32, 8, 9, 9, 9, 11]), 1):
	c = wsum.cell(row=4, column=i, value=h)
	c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
	c.fill = HEAD_FILL
	c.border = BORDER
	wsum.column_dimensions[get_column_letter(i)].width = w
GR, SR = f"Checklist!$B$2:$B${last + 100}", f"Checklist!$H$2:$H${last + 100}"
row = 4
for g in groups + ["All"]:
	row += 1
	is_all = g == "All"
	cell(wsum, row, 1, g, bold=is_all)
	cell(wsum, row, 2, f"=COUNTIF({GR},$A{row})" if not is_all else f"=SUM(B5:B{row - 1})", bold=is_all)
	for j, st in enumerate(STATUSES, 3):
		col = get_column_letter(j)
		cell(wsum, row, j, f"=COUNTIFS({GR},$A{row},{SR},{col}$4)" if not is_all else f"=SUM({col}5:{col}{row - 1})", bold=is_all)

# ---------------------------------------------------------------- Change Log
wl = wb.create_sheet("Change Log")
header(wl, ["Date", "Area", "Items", "What changed", "By"], [13, 20, 34, 100, 16])
wl.freeze_panes = "A2"
for i, row in enumerate(CHANGELOG, 2):
	for j, v in enumerate(row, 1):
		cell(wl, i, j, v, wrap=j in (3, 4))

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "items", len(CHECKLIST))
