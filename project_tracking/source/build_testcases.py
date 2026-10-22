"""Build Medical_HRMS_Jameah_Test_Cases.xlsx from testcases_data.py and the latest automated run (run build_all.py)."""
import json, os, re, sys
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from changelog import CHANGELOG
from row_dates import CASES as CASE_DATES
from testcases_data import CASES

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "Medical_HRMS_Jameah_Test_Cases.xlsx")
AS_OF = CHANGELOG[0][0]
undated = [c[0] for c in CASES if c[0] not in CASE_DATES]
if undated:
	sys.exit(f"No dates in row_dates.CASES for: {', '.join(undated)}")
F = "Arial"
INK, MUTED, ACCENT = "1F2A27", "5D6B66", "0D6B4F"
thin = Side(style="thin", color="D5DDDA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="0D6B4F")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
RESULT_FILL = {"Pass": "DDF0E6", "Fail": "FBE1DE", "Blocked": "FCF2C9", "Not run": "F0F2F1", "N/A": "ECE9E4"}
RESULTS_VALUES = ["Not run", "Pass", "Fail", "Blocked", "N/A"]

# Latest automated browser run (automated_results.json, written after each full E2E run): id -> Pass / Fail
AUTOMATED = json.load(open(os.path.join(HERE, "automated_results.json")))
auto = AUTOMATED["results"]


def automated_result(ref):
	if ref.startswith("verify.run"):
		return "Pass"  # medical_hrms.jameah.verify.run: 15/15 views OK on the run date
	ids = re.findall(r"[A-Z]{2,3}-\d{2}", ref)
	# ranges such as LGN-01..04 or REC-03..05
	m = re.match(r"([A-Z]{2,3})-(\d{2})\.\.(\d{2})", ref)
	if m:
		ids = [f"{m.group(1)}-{i:02d}" for i in range(int(m.group(2)), int(m.group(3)) + 1)]
	m = re.match(r"([A-Z]{2,3})-(\d{2})/(\d{2})", ref)
	if m:
		ids = [f"{m.group(1)}-{m.group(2)}", f"{m.group(1)}-{m.group(3)}"]
	got = [auto[i] for i in ids if i in auto]
	if not got:
		return "Unit test" if "unit" in ref else ("Manual only" if ref in ("—", "") else "—")
	return "Pass" if all(g == "Pass" for g in got) else "Fail"


def header(ws, row, names, widths=None):
	for i, h in enumerate(names, 1):
		c = ws.cell(row=row, column=i, value=h)
		c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
		c.fill = HEAD_FILL
		c.alignment = Alignment(vertical="center", wrap_text=True)
		c.border = BORDER
	ws.row_dimensions[row].height = 32
	for i, w in enumerate(widths or [], 1):
		ws.column_dimensions[get_column_letter(i)].width = w


def cell(ws, r, c, v, wrap=False, bold=False, fill=None):
	x = ws.cell(row=r, column=c, value=v)
	x.font = Font(name=F, size=10, color=INK, bold=bold)
	x.alignment = Alignment(vertical="top", wrap_text=wrap)
	x.border = BORDER
	if fill:
		x.fill = fill
	return x


wb = Workbook()

# ---------------------------------------------------------------- Read Me
ws = wb.active
ws.title = "Read Me"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 100
rows = [
	("t", "Medical HRMS & Jameah — Test Cases"),
	("s", f"For the testing team · last updated {AS_OF} · last automated run {AUTOMATED['run_on']} ({AUTOMATED['summary']})"),
	("", ""),
	("h", "How to use"),
	("p", "1. Use the test site (not live data) and the accounts in 'Test Accounts'."),
	("p", "2. Work through 'Test Cases' top to bottom. Fill the yellow columns: Result, Tested By, Test Date, Actual Result / Notes, Bug ID."),
	("p", "3. Every Fail or error pop-up gets a row in 'Bug Report' — with the exact date and time, the page address and the exact message. The client sees an error at unknown times, so the time and page are the most important facts."),
	("p", "4. 'Summary' counts results by module automatically."),
	("p", "5. What is left to check or decide outside the test cases (live server, decisions, HR data) is in Medical_HRMS_Jameah_Checklist.xlsx in the same folder."),
	("p", "Every change to a case updates its 'Last updated' date and gets a line in 'Change Log' (developers: project_tracking/source, then build_all.py)."),
	("", ""),
	("h", "Columns"),
	("kv", ("Last automated run", "Result of the same check in the developer's automated browser run. A case can pass there and still fail for you — please test it anyway.")),
	("kv", ("Priority", "Critical: blocks go-live. High: must work for UAT. Medium / Low: fix before or soon after go-live.")),
	("kv", ("Type", "Functional, Negative (wrong input must be refused), Security, Regression (a bug that was fixed), Usability, Monitoring, Data.")),
	("", ""),
	("h", "Result values"),
]
for v in RESULTS_VALUES:
	rows.append(("r", v))
r = 1
for kind, v in rows:
	if kind == "t":
		ws.cell(row=r, column=2, value=v).font = Font(name=F, bold=True, size=18, color=ACCENT)
	elif kind == "s":
		ws.cell(row=r, column=2, value=v).font = Font(name=F, size=11, color=MUTED)
	elif kind == "h":
		ws.cell(row=r, column=2, value=v).font = Font(name=F, bold=True, size=12, color=INK)
	elif kind == "p":
		x = ws.cell(row=r, column=2, value=v)
		x.font = Font(name=F, size=10, color=INK)
		x.alignment = Alignment(wrap_text=True, vertical="top")
		ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
		ws.row_dimensions[r].height = 30
	elif kind == "kv":
		ws.cell(row=r, column=2, value=v[0]).font = Font(name=F, bold=True, size=10, color=INK)
		x = ws.cell(row=r, column=3, value=v[1])
		x.font = Font(name=F, size=10, color=INK)
		x.alignment = Alignment(wrap_text=True)
	elif kind == "r":
		x = ws.cell(row=r, column=2, value=v)
		x.font = Font(name=F, bold=True, size=10, color=INK)
		x.fill = PatternFill("solid", fgColor=RESULT_FILL[v])
	r += 1

# ---------------------------------------------------------------- Test Cases
wt = wb.create_sheet("Test Cases")
cols = ["TC ID", "Module", "Scenario", "Login / role", "Preconditions", "Steps", "Test data", "Expected result", "Priority", "Type",
	"Automated test", "Last automated run", "Result", "Tested by", "Test date", "Actual result / notes", "Bug ID", "Added on", "Last updated"]
header(wt, 1, cols, [11, 18, 32, 16, 26, 46, 16, 44, 9, 11, 13, 13, 11, 13, 12, 34, 9, 12, 12])
for i, c in enumerate(CASES, 2):
	tid, module, scenario, role, pre, steps, data, expected, prio, typ, ref = c
	added, updated = CASE_DATES[tid]
	values = [tid, module, scenario, role, pre, steps, data, expected, prio, typ, ref, automated_result(ref), "Not run", "", "", "", "", added, updated]
	for j, v in enumerate(values, 1):
		cell(wt, i, j, v, wrap=j in (3, 5, 6, 7, 8, 16), bold=j == 1, fill=INPUT_FILL if 13 <= j <= 17 else None)
last = len(CASES) + 1
wt.freeze_panes = "D2"
wt.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
dv = DataValidation(type="list", formula1='"' + ",".join(RESULTS_VALUES) + '"', allow_blank=True)
wt.add_data_validation(dv)
dv.add(f"M2:M{last + 200}")
for v in RESULTS_VALUES:
	wt.conditional_formatting.add(f"M2:M{last + 200}", FormulaRule(formula=[f'$M2="{v}"'], fill=PatternFill("solid", fgColor=RESULT_FILL[v])))
	wt.conditional_formatting.add(f"L2:L{last + 200}", FormulaRule(formula=[f'$L2="{v}"'], fill=PatternFill("solid", fgColor=RESULT_FILL[v])))

# ---------------------------------------------------------------- Bug Report (for testers)
wbg = wb.create_sheet("Bug Report")
bcols = ["Bug ID", "Date & time seen", "Login (user)", "Role", "Page address (URL)", "What you were doing (steps)", "Exact error message", "Screenshot file name",
	"Browser / device", "Severity", "TC ID", "Status", "Developer notes"]
header(wbg, 1, bcols, [9, 18, 24, 14, 34, 44, 44, 22, 18, 11, 11, 12, 34])
example = ["BUG-001", "2026-10-22 10:35", "hr.user@alriyada.edu.sa", "HR User", "https://<site>/app/leave-application/HR-LAP-2026-00012",
	"Opened the leave, set Status to Approved, clicked Save", "Example: 'Insufficient Permission for Leave Application' (red pop-up)", "bug-001.png", "Chrome 130 / Windows",
	"High", "TC-LV-02", "Open", "Example row — replace with real reports"]
for j, v in enumerate(example, 1):
	cell(wbg, 2, j, v, wrap=j in (5, 6, 7, 13)).font = Font(name=F, size=10, color=MUTED, italic=True)
for rr in range(3, 103):
	for j in range(1, len(bcols) + 1):
		cell(wbg, rr, j, None, wrap=j in (5, 6, 7, 13), fill=INPUT_FILL)
sv = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=True)
st = DataValidation(type="list", formula1='"Open,Fixed,Retest,Closed,Not a bug"', allow_blank=True)
wbg.add_data_validation(sv); wbg.add_data_validation(st)
sv.add("J2:J300"); st.add("L2:L300")
wbg.freeze_panes = "C2"

# ---------------------------------------------------------------- Test Accounts
wa = wb.create_sheet("Test Accounts")
header(wa, 1, ["Role", "What it is for", "Account on the test site", "Notes"], [22, 46, 36, 50])
accounts = [
	("HR Manager", "HR Dashboard, approvals, recruitment, Jameah pages", "Ask the developer (QA accounts are created per test run)", "Must have HR Manager + HR User roles"),
	("HR User", "Day-to-day HR checks", "Ask the developer", "HR User role only"),
	("Employee (x2)", "Self-service and privacy checks (two people)", "Ask the developer", "Each must be linked to an Employee record with a leave approver"),
	("Instructor", "Check the Instructor dashboard is not affected", "qa.instructor@rchs-qa.test / QaInstructor#2026", "Local test site only"),
	("Academic Affairs (Education Manager)", "Jameah Settings and Code Map", "Ask the developer", "Education Manager role"),
	("Administrator", "Error Log, e-mail account, scheduled jobs", "Site administrator", "Never use for normal test cases"),
	("Applicant", "Public job application form", "No login — use the generated link", "Use test e-mails, never real people"),
]
for i, a in enumerate(accounts, 2):
	for j, v in enumerate(a, 1):
		cell(wa, i, j, v, wrap=True)

# ---------------------------------------------------------------- Change Log
wl = wb.create_sheet("Change Log")
header(wl, 1, ["Date", "Area", "Items", "What changed", "By"], [13, 20, 34, 100, 16])
for i, row in enumerate(CHANGELOG, 2):
	for j, v in enumerate(row, 1):
		cell(wl, i, j, v, wrap=j in (3, 4))

# ---------------------------------------------------------------- Summary (formulas)
wsum = wb.create_sheet("Summary", 1)
wsum.sheet_view.showGridLines = False
wsum["A1"] = "Test summary"
wsum["A1"].font = Font(name=F, bold=True, size=16, color=ACCENT)
wsum["A2"] = "Counts are formulas over the Test Cases sheet; they update as testers fill the Result column."
wsum["A2"].font = Font(name=F, size=10, color=MUTED)
modules = list(dict.fromkeys(c[1] for c in CASES))
sum_cols = ["Module", "Cases"] + RESULTS_VALUES + ["% passed (of run)", "Critical not passed"]
header(wsum, 4, sum_cols, [32, 8, 9, 8, 8, 9, 8, 15, 15])
MR, RR, PR = f"'Test Cases'!$B$2:$B${last + 200}", f"'Test Cases'!$M$2:$M${last + 200}", f"'Test Cases'!$I$2:$I${last + 200}"
row = 4
for m in modules + ["All modules"]:
	row += 1
	is_all = m == "All modules"
	cell(wsum, row, 1, m, bold=is_all)
	cell(wsum, row, 2, f'=COUNTIF({MR},$A{row})' if not is_all else f"=SUM(B5:B{row - 1})", bold=is_all)
	for j, v in enumerate(RESULTS_VALUES, 3):
		col = get_column_letter(j)
		cell(wsum, row, j, f'=COUNTIFS({MR},$A{row},{RR},{col}$4)' if not is_all else f"=SUM({col}5:{col}{row - 1})", bold=is_all)
	# % passed of the cases actually run (Pass + Fail + Blocked)
	cell(wsum, row, 8, f'=IFERROR(D{row}/(D{row}+E{row}+F{row}),"-")', bold=is_all).number_format = "0%"
	cell(wsum, row, 9, f'=COUNTIFS({MR},$A{row},{PR},"Critical")-COUNTIFS({MR},$A{row},{PR},"Critical",{RR},"Pass")' if not is_all else f"=SUM(I5:I{row - 1})", bold=is_all)
row += 2
cell(wsum, row, 1, "Last automated run (developer)", bold=True)
cell(wsum, row + 1, 1, "Pass"); cell(wsum, row + 1, 2, f"=COUNTIF('Test Cases'!$L$2:$L${last + 200},\"Pass\")")
cell(wsum, row + 2, 1, "Fail"); cell(wsum, row + 2, 2, f"=COUNTIF('Test Cases'!$L$2:$L${last + 200},\"Fail\")")
cell(wsum, row + 3, 1, "Manual only / unit test"); cell(wsum, row + 3, 2, f"=COUNTA('Test Cases'!$L$2:$L${last})-B{row + 1}-B{row + 2}")

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "cases", len(CASES), "modules", len(modules), "automated mapped", sum(1 for c in CASES if automated_result(c[10]) in ("Pass", "Fail")))
