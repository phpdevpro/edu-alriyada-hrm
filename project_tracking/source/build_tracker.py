"""Build Medical_HRMS_Jameah_Task_Tracker.xlsx from tracker_data.py (run build_all.py)."""
import os
import sys
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from changelog import CHANGELOG
from row_dates import TRACKER as UPDATED
from tracker_data import ASSIGNED_TO, AS_OF, BUGS, DECISIONS, STATUSES, TASKS, TEAM
from colours import BUG_STATUS, PRIORITY, TASK_STATUS, badge, colour_column

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "Medical_HRMS_Jameah_Task_Tracker.xlsx")
AS_OF = CHANGELOG[0][0] if CHANGELOG else AS_OF
DECISIONS = sorted(DECISIONS, key=lambda d: int(d[0].split("-")[1]))
undated = [r[0] for r in [*TASKS, *BUGS, *DECISIONS] if r[0] not in UPDATED]
if undated:
	sys.exit(f"No 'last updated' date in row_dates.TRACKER for: {', '.join(undated)}")
F = "Arial"
INK, MUTED, ACCENT = "1F2A27", "5D6B66", "0D6B4F"
thin = Side(style="thin", color="D5DDDA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="0D6B4F")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")


def style_header(ws, row, ncols):
	for c in range(1, ncols + 1):
		cell = ws.cell(row=row, column=c)
		cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
		cell.fill = HEAD_FILL
		cell.alignment = Alignment(vertical="center", wrap_text=True)
		cell.border = BORDER


def table(ws, headers, rows, widths, start_row=1, wrap_cols=()):
	for i, h in enumerate(headers, 1):
		ws.cell(row=start_row, column=i, value=h)
	style_header(ws, start_row, len(headers))
	ws.row_dimensions[start_row].height = 30
	for r, row in enumerate(rows, start_row + 1):
		for c, value in enumerate(row, 1):
			cell = ws.cell(row=r, column=c, value=value)
			cell.font = Font(name=F, size=10, color=INK)
			cell.alignment = Alignment(vertical="top", wrap_text=c in wrap_cols)
			cell.border = BORDER
	for i, w in enumerate(widths, 1):
		ws.column_dimensions[get_column_letter(i)].width = w
	ws.freeze_panes = ws.cell(row=start_row + 1, column=3)
	ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
	return start_row + len(rows)


wb = Workbook()

# ------------------------------------------------------------------ Read Me
ws = wb.active
ws.title = "Read Me"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 100
lines = [
	("title", "Medical HRMS & Jameah — Task Tracker"),
	("sub", f"Al-Riyada College for Health Sciences · last updated {AS_OF}"),
	("", ""),
	("h", "What this workbook is"),
	("p", "One place to see what has been done on Medical HRMS (HR system) and Jameah (the Ministry data system), what is being done now, and what is still pending."),
	("p", "It continues the 1 Oct 2026 tracker (HRM_TASK_TRACKER.md / .csv): tasks HR-01 to HR-32 keep their IDs and are updated; HR-33 onwards and JM-01 onwards are new."),
	("", ""),
	("h", "Sheets"),
	("kv", ("Summary", "Counts by status, workstream and date. Updates by itself when you change a status in Tasks.")),
	("kv", ("Tasks", "Every task: when it was done, status, owner (who leads it: Abdul Hakeem for all development), assigned to (who the task is given to), evidence and the next step. Filter on any column.")),
	("kv", ("Bugs", "Every bug found during testing, its cause, the fix and how it was verified.")),
	("kv", ("Decisions", "Questions only the college, the Ministry or another team can answer. Yellow columns are for the answer.")),
	("kv", ("Change Log", "Every change, newest first, with its date and the items it touched.")),
	("", ""),
	("h", "Status values"),
]
for s in STATUSES:
	meaning = {
		"Done": "Built and tested.",
		"Done – client UAT pending": "Built and tested by the team; waiting for HR/client acceptance testing.",
		"In progress": "Being worked on.",
		"Pending – decision needed": "Waiting for an answer (see Decisions).",
		"Pending – data / setup": "Code is ready; needs data entry or server setup.",
		"Pending – development": "Not built yet.",
		"Pending – other team": "Needs Finance, the server admin or another party outside the HRMS team.",
		"Planned": "Agreed for later; not started.",
	}[s]
	lines.append(("status", (s, meaning)))
lines += [("", ""), ("h", "Priority")]
for p, meaning in (("Critical", "Blocks go-live or the Ministry data: do first."), ("High", "Needed for UAT / go-live."),
		("Medium", "Needed; can follow go-live."), ("Low", "Tidy-up or nice to have.")):
	lines.append(("prio", (p, meaning)))
lines += [("", ""), ("h", "How to update"),
	("p", "Every change: update the row (status, evidence, next step) and its 'Last updated' date, and add a dated line to the Change Log. The same goes for the Checklist and the Test Cases workbook."),
	("p", "Developers: edit project_tracking/source/*.py (data, dates in row_dates.py, changelog.py) and run python3 project_tracking/source/build_all.py - all three workbooks are rebuilt."),
	("", ""), ("h", "Sources"),
	("p", "HRM_TASK_TRACKER.md/.csv, HRM_UAT_SCENARIOS.md, HRM_APPROVAL_MATRIX.md, JAMEAH_HRMS_* documents, MINISTRY_FIELD_MAPPING.html, Ministry schema V0.1, git history of medical_hrms (phpdevpro/edu-alriyada-hrm), test runs of 19–21 Oct 2026."),
	("p", "Companion files in the same project_tracking folder: Medical_HRMS_Jameah_Checklist.xlsx (what is left to check or decide) and Medical_HRMS_Jameah_Test_Cases.xlsx (for the testing team).")]
r = 1
for kind, value in lines:
	if kind == "title":
		ws.cell(row=r, column=2, value=value).font = Font(name=F, bold=True, size=18, color=ACCENT)
	elif kind == "sub":
		ws.cell(row=r, column=2, value=value).font = Font(name=F, size=11, color=MUTED)
	elif kind == "h":
		ws.cell(row=r, column=2, value=value).font = Font(name=F, bold=True, size=12, color=INK)
	elif kind == "p":
		c = ws.cell(row=r, column=2, value=value)
		c.font = Font(name=F, size=10, color=INK)
		c.alignment = Alignment(wrap_text=True, vertical="top")
		ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
		ws.row_dimensions[r].height = 30
	elif kind in ("kv", "status", "prio"):
		a, b = value
		ca = ws.cell(row=r, column=2, value=a)
		ca.font = Font(name=F, bold=True, size=10, color=INK)
		if kind == "status":
			badge(ca, TASK_STATUS, a, F)
		elif kind == "prio":
			badge(ca, PRIORITY, a, F)
		cb = ws.cell(row=r, column=3, value=b)
		cb.font = Font(name=F, size=10, color=INK)
		cb.alignment = Alignment(wrap_text=True)
	r += 1

# ------------------------------------------------------------------ Tasks
wt = wb.create_sheet("Tasks")
headers = ["ID", "Workstream", "Task", "When done", "Status", "Owner", "Assigned to", "Priority", "Done so far / evidence", "Next step / pending", "Reference (commit / file)", "Last updated"]
assigned = lambda t: ASSIGNED_TO.get(t[0]) or ("Abdul Hakeem" if "Abdul Hakeem" in t[5] else t[5])
last = table(wt, headers, [t[:6] + (assigned(t),) + t[6:] + (UPDATED[t[0]],) for t in TASKS], [8, 16, 42, 16, 22, 20, 16, 9, 60, 50, 28, 13], wrap_cols=(3, 9, 10, 11))
n_tasks = len(TASKS)
dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
dv.error, dv.errorTitle = "Pick a status from the list.", "Status"
wt.add_data_validation(dv)
dv.add(f"E2:E{last + 200}")
pv = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=True)
wt.add_data_validation(pv)
pv.add(f"H2:H{last + 200}")
# the lead picks who a task is given to; other names can still be typed (no error message)
av = DataValidation(type="list", formula1='"' + ",".join(TEAM) + '"', allow_blank=True, showErrorMessage=False)
wt.add_data_validation(av)
av.add(f"G2:G{last + 200}")
colour_column(wt, "E", 2, last + 200, TASK_STATUS)
colour_column(wt, "H", 2, last + 200, PRIORITY)
for row in range(2, last + 1):
	wt.cell(row=row, column=1).font = Font(name=F, bold=True, size=10, color=INK)

# ------------------------------------------------------------------ Bugs
wbg = wb.create_sheet("Bugs")
blast = table(wbg, ["Bug ID", "Found on", "Area", "Problem", "Cause", "Fix", "Commit / task", "Status", "How verified", "Last updated"],
	[b + (UPDATED[b[0]],) for b in BUGS], [8, 13, 20, 46, 42, 40, 13, 18, 20, 13], wrap_cols=(4, 5, 6, 9))
colour_column(wbg, "H", 2, blast + 100, BUG_STATUS)
bv = DataValidation(type="list", formula1='"Fixed,Open,Open – other team,Open – check server"', allow_blank=True)
wbg.add_data_validation(bv)
bv.add(f"H2:H{blast + 100}")

# ------------------------------------------------------------------ Decisions
wd = wb.create_sheet("Decisions")
dlast = table(wd, ["ID", "Decision needed", "Who decides", "Options / recommendation", "Blocks tasks", "Answer (fill in)", "Answered on", "Last updated"],
	[d[:5] + (d[5:] + ("", ""))[:2] + (UPDATED[d[0]],) for d in DECISIONS], [9, 52, 22, 46, 18, 34, 13, 13], wrap_cols=(2, 4, 6))
for row in range(2, dlast + 1):
	for col in (6, 7):
		wd.cell(row=row, column=col).fill = INPUT_FILL

# ------------------------------------------------------------------ Change Log
wl = wb.create_sheet("Change Log")
table(wl, ["Date", "Area", "Items", "What changed", "By"], CHANGELOG, [13, 20, 34, 100, 16], wrap_cols=(3, 4))

# ------------------------------------------------------------------ Summary (formulas over Tasks)
wsum = wb.create_sheet("Summary", 1)
wsum.sheet_view.showGridLines = False
wsum.column_dimensions["A"].width = 30
wsum["A1"] = "Summary"
wsum["A1"].font = Font(name=F, bold=True, size=16, color=ACCENT)
wsum["A2"] = f"Counts are formulas over the Tasks and Bugs sheets. As of {AS_OF}."
wsum["A2"].font = Font(name=F, size=10, color=MUTED)
TR = f"Tasks!$E$2:$E${last + 200}"
WR = f"Tasks!$B$2:$B${last + 200}"
DR = f"Tasks!$D$2:$D${last + 200}"
IR = f"Tasks!$A$2:$A${last + 200}"
row = 4
wsum.cell(row=row, column=1, value="Tasks by status")
wsum.cell(row=row, column=2, value="Tasks")
style_header(wsum, row, 2)
for s in STATUSES:
	row += 1
	a = wsum.cell(row=row, column=1, value=s)
	badge(a, TASK_STATUS, s, F)
	wsum.cell(row=row, column=2, value=f'=COUNTIF({TR},A{row})')
first_status_row, last_status_row = 5, row
row += 1
wsum.cell(row=row, column=1, value="Total tasks").font = Font(name=F, bold=True, size=10)
wsum.cell(row=row, column=2, value=f"=SUM(B{first_status_row}:B{last_status_row})").font = Font(name=F, bold=True, size=10)
total_row = row

row += 2
wsum.cell(row=row, column=1, value="Medical HRMS vs Jameah")
for i, h in enumerate(["Done (both Done statuses)", "Open (everything else)", "Total"], 2):
	wsum.cell(row=row, column=i, value=h)
style_header(wsum, row, 4)
for label, prefix in (("Medical HRMS (HR-…)", "HR-*"), ("Jameah (JM-…)", "JM-*")):
	row += 1
	wsum.cell(row=row, column=1, value=label)
	wsum.cell(row=row, column=2, value=f'=COUNTIFS({IR},"{prefix}",{TR},"Done")+COUNTIFS({IR},"{prefix}",{TR},"Done – client UAT pending")')
	wsum.cell(row=row, column=3, value=f'=D{row}-B{row}')
	wsum.cell(row=row, column=4, value=f'=COUNTIF({IR},"{prefix}")')

row += 2
wsum.cell(row=row, column=1, value="Work done by date")
wsum.cell(row=row, column=2, value="Tasks")
style_header(wsum, row, 2)
for d in list(dict.fromkeys(t[3] for t in TASKS if t[3] != "—")) + ["—"]:
	row += 1
	wsum.cell(row=row, column=1, value=d if d != "—" else "— (not done yet)")
	wsum.cell(row=row, column=2, value=f'=COUNTIF({DR},"{d}")')

row += 2
workstreams = list(dict.fromkeys(t[1] for t in TASKS))
wsum.cell(row=row, column=1, value="Workstream")
for i, s in enumerate(STATUSES, 2):
	wsum.cell(row=row, column=i, value=s)
	wsum.column_dimensions[get_column_letter(i)].width = 15
wsum.cell(row=row, column=len(STATUSES) + 2, value="Total")
style_header(wsum, row, len(STATUSES) + 2)
wsum.row_dimensions[row].height = 45
head_row = row
for w in workstreams:
	row += 1
	wsum.cell(row=row, column=1, value=w)
	for i, s in enumerate(STATUSES, 2):
		col = get_column_letter(i)
		wsum.cell(row=row, column=i, value=f'=COUNTIFS({WR},$A{row},{TR},{col}${head_row})')
	wsum.cell(row=row, column=len(STATUSES) + 2, value=f'=COUNTIF({WR},$A{row})')

row += 2
wsum.cell(row=row, column=1, value="Bugs")
wsum.cell(row=row, column=2, value="Count")
style_header(wsum, row, 2)
for s in ("Fixed", "Open", "Open – other team", "Open – check server"):
	row += 1
	badge(wsum.cell(row=row, column=1, value=s), BUG_STATUS, s, F)
	wsum.cell(row=row, column=2, value=f'=COUNTIF(Bugs!$H$2:$H${blast + 100},A{row})')

for rr in wsum.iter_rows(min_row=4, max_row=row):
	for c in rr:
		if c.value is None:
			continue
		c.border = BORDER
		if c.fill.fgColor.rgb in ("000D6B4F", "FF0D6B4F"):
			continue  # header cell, keep white bold text
		c.font = Font(name=F, size=10, color=INK, bold=bool(c.font and c.font.bold))

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "tasks", n_tasks, "bugs", len(BUGS), "decisions", len(DECISIONS))
