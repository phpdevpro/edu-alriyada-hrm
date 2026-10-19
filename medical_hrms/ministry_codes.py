"""Ministry lookup codes: one dataset, kept in sync on every migrate.

Lookup fields (MINISTRY_FIELD_MAPPING.html "Lookup" data type) are select-
only: their values must be codes from the Ministry workbook, stored as
Jameah Ministry Code records named "<category>::<code>". The Ministry
receives the code; people pick by the Arabic / English name.

Source: two editions of the Ministry workbook with identical codes -
"Ministry Code - (Arabic).xlsx" (real Arabic names) and
"Ministry Code - (Translated).xlsx" (English sheet names and translated
names). build_dataset() merges them into data/ministry_codes.csv, which is
committed; sync() loads that CSV on every migrate (skipped when unchanged).

Rebuild the CSV after the Ministry sends a new workbook:
    bench --site <site> execute medical_hrms.ministry_codes.build_dataset \
        --kwargs "{'arabic': '/path/Ministry Code - (Arabic).xlsx', 'translated': '/path/Ministry Code - (Translated).xlsx'}"
"""
import csv
import hashlib
import os
import re

import frappe

DOCTYPE = "Jameah Ministry Code"
DATASET = os.path.join(os.path.dirname(__file__), "data", "ministry_codes.csv")
HASH_KEY = "medical_hrms_ministry_codes_hash"
COLUMNS = ("code_category", "ministry_code", "name_english", "name_arabic")

# Arabic header cells use decorative stretching (tatweel), e.g. "الــرمــز".
CODE_HEADERS = {"الرمز", "الترميز"}
ARABIC_NAME_HEADERS = {"الاسم", "الإسم"}
ENGLISH_NAME_HEADERS = {"الاسم بالإنجليزي", "الاسم بالانجليزي"}


def _text(value):
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	return re.sub(r"\s+", " ", str(value)).strip()


def _arabic_header(value):
	return re.sub(r"\s+", " ", re.sub(r"[ـ\-]+", "", _text(value))).strip()


def _sheet_number(title):
	match = re.search(r"\d+", title)
	return match.group(0) if match else None


def _sheet_rows(sheet):
	"""Non-empty rows; stops after a long blank run (one sheet is 1M rows)."""
	blank = 0
	for row in sheet.iter_rows(values_only=True):
		if all(_text(v) == "" for v in row):
			blank += 1
			if blank > 50:
				return
			continue
		blank = 0
		yield row


def _translated_codes(rows, header):
	"""(code, english) per row of the Translated edition, using the existing
	importer's header detection and code rules so categories/codes match."""
	from medical_hrms import import_ministry_excel_to_jameah_codes as legacy

	out = {}
	for row in rows[header["header_row"]:]:
		cell = lambda key: row[header[key]] if header[key] is not None and len(row) > header[key] else None
		code, _source = legacy._pick_code(cell("symbol_col"), cell("coding_col"))
		if not code or code in out:
			continue
		# Prefer the Ministry's own English name; otherwise the translated
		# "Name" column (this edition translated the Arabic names to English).
		out[code] = _text(cell("name_en_col")) or _text(cell("name_ar_col"))
	return out


def _arabic_codes(rows):
	header = [_arabic_header(v) for v in rows[0]]
	pick = lambda names: next((i for i, h in enumerate(header) if h in names), None)
	code_col, ar_col, en_col = pick(CODE_HEADERS), pick(ARABIC_NAME_HEADERS), pick(ENGLISH_NAME_HEADERS)
	out = {}
	for row in rows[1:]:
		code = _text(row[code_col]).upper() if code_col is not None and len(row) > code_col else ""
		if code and code not in out:
			out[code] = (
				_text(row[ar_col]) if ar_col is not None and len(row) > ar_col else "",
				_text(row[en_col]) if en_col is not None and len(row) > en_col else "",
			)
	return out


def _header_map(rows):
	"""The legacy importer's header detection, on already-read rows."""
	from medical_hrms import import_ministry_excel_to_jameah_codes as legacy

	best = None
	for index, row in enumerate(rows[:12], start=1):
		tokens = [legacy._norm_header_token(v) for v in row]
		found = {"header_row": index, "symbol_col": None, "coding_col": None, "name_ar_col": None, "name_en_col": None}
		for col, token in enumerate(tokens):
			if not token:
				continue
			if found["symbol_col"] is None and ("symbol" in token or "الرمز" in token):
				found["symbol_col"] = col
			if found["coding_col"] is None and "coding" in token:
				found["coding_col"] = col
			if found["name_en_col"] is None and "name in english" in token:
				found["name_en_col"] = col
			if found["name_ar_col"] is None and (token in ("name", "the name") or "name in arabic" in token):
				found["name_ar_col"] = col
		score = sum(v is not None for key, v in found.items() if key != "header_row")
		if score and (not best or score > best[0]):
			best = (score, found)
	return best[1] if best else {"header_row": 1, "symbol_col": 0, "coding_col": None, "name_ar_col": 1, "name_en_col": 2}


def build_dataset(arabic, translated, output=DATASET):
	"""Merge both workbook editions into the committed CSV."""
	from openpyxl import load_workbook

	from medical_hrms import import_ministry_excel_to_jameah_codes as legacy

	ar_book = load_workbook(arabic, read_only=True, data_only=True)
	tr_book = load_workbook(translated, read_only=True, data_only=True)
	ar_sheets = {_sheet_number(ws.title): ws for ws in ar_book.worksheets}

	records, report = [], []
	for tr_sheet in tr_book.worksheets:
		number = _sheet_number(tr_sheet.title)
		if number not in ar_sheets:
			frappe.throw(f"Arabic workbook has no sheet {number} ({tr_sheet.title.strip()}).")
		category = legacy._normalize_category(tr_sheet.title)
		tr_rows = list(_sheet_rows(tr_sheet))
		english = _translated_codes(tr_rows, _header_map(tr_rows))
		arabic_names = _arabic_codes(list(_sheet_rows(ar_sheets[number])))

		missing_arabic = 0
		for code, name_english in english.items():
			name_arabic, original_english = arabic_names.get(code, ("", ""))
			missing_arabic += not name_arabic
			records.append((category, code, original_english or name_english or name_arabic or code,
				name_arabic or name_english or code))
		extra = sorted(set(arabic_names) - set(english))
		report.append((category, len(english), missing_arabic, extra))

	os.makedirs(os.path.dirname(output), exist_ok=True)
	with open(output, "w", newline="", encoding="utf-8") as handle:
		writer = csv.writer(handle)
		writer.writerow(COLUMNS)
		writer.writerows(records)

	print(f"Wrote {len(records)} codes to {output}")
	for category, count, missing_arabic, extra in report:
		note = f" | {missing_arabic} without Arabic name" if missing_arabic else ""
		note += f" | only in Arabic edition: {extra[:5]}" if extra else ""
		print(f"  {category:40} {count:5}{note}")
	return len(records)


def _load_dataset():
	with open(DATASET, encoding="utf-8") as handle:
		content = handle.read()
	rows = list(csv.DictReader(content.splitlines()))
	return rows, hashlib.sha256(content.encode("utf-8")).hexdigest()


def sync(force=False):
	"""after_install / after_migrate: make the table match the dataset.
	Inserts new codes and updates changed names. Never deletes (codes may be
	referenced); see remove_non_ministry_codes() for that."""
	if not frappe.db.exists("DocType", DOCTYPE) or not os.path.exists(DATASET):
		return
	rows, digest = _load_dataset()
	if not int(force) and frappe.db.get_global(HASH_KEY) == digest:
		return

	existing = {
		(r.code_category, r.ministry_code): r
		for r in frappe.get_all(DOCTYPE, fields=["name", *COLUMNS], limit_page_length=0)
	}
	inserted = updated = 0
	for row in rows:
		current = existing.get((row["code_category"], row["ministry_code"]))
		if not current:
			frappe.get_doc({"doctype": DOCTYPE, **row}).insert(ignore_permissions=True)
			inserted += 1
		elif current.name_english != row["name_english"] or current.name_arabic != row["name_arabic"]:
			frappe.db.set_value(DOCTYPE, current.name, {
				"name_english": row["name_english"], "name_arabic": row["name_arabic"]}, update_modified=False)
			updated += 1

	frappe.db.set_global(HASH_KEY, digest)
	frappe.db.commit()
	frappe.clear_cache(doctype=DOCTYPE)
	print(f"Ministry codes synced: {inserted} inserted, {updated} updated, {len(rows)} in dataset.")


# Demo seed codes (seed_data/jameah_dummy_data.json) that records still use,
# mapped to the real Ministry code with the same meaning. Unreferenced demo
# codes need no entry; they are simply deleted.
DEMO_REPLACEMENTS = {
	"Nationality::N-SA": "Nationality::101",  # Saudi Arabia
	"Nationality::N-EG": "Nationality::118",  # Egypt
	"Gender::G-MALE": "Gender::1",
	"Gender::G-FEMALE": "Gender::2",
	"Identity type::ID-NID": "Identity type::01",  # National ID
	"Identity type::ID-IQ": "Identity type::03",  # Iqama
	"Coding cities and governora::CITY-RIYADH": "Coding cities and governora::0101000",  # Riyadh City
	"Coding cities and governora::CITY-JEDDAH": "Coding cities and governora::0202000",  # Jeddah Governorate
	"Degree::DEG-BSC": "Coding of academic degrees::2",  # Bachelors degree
	"Degree::DEG-MSC": "Coding of academic degrees::4",  # Masters degree
	"Degree::DEG-PHD": "Coding of academic degrees::5",  # Doctorate / PhD
	"Specialization::SPEC-MED": "Specialization Coding Guide::HA09",  # medicine
}


def lookup_fields():
	"""Every (doctype, fieldname) linking to Jameah Ministry Code, whether
	defined on the DocType, as a Custom Field, or re-pointed by a Property
	Setter (e.g. Employee.gender)."""
	return frappe.db.sql("""
		select parent, fieldname from `tabDocField` where fieldtype='Link' and options=%(dt)s
		union select dt, fieldname from `tabCustom Field` where fieldtype='Link' and options=%(dt)s
		union select doc_type, field_name from `tabProperty Setter` where property='options' and value=%(dt)s
	""", {"dt": DOCTYPE})


def remove_non_ministry_codes(commit=0):
	"""Replace references to codes that are not in the Ministry dataset, then
	delete those codes. Dry run unless commit=1. Stops without changing
	anything if a referenced code has no replacement."""
	commit = int(commit)
	rows, _digest = _load_dataset()
	ministry = {f"{r['code_category']}::{r['ministry_code']}" for r in rows}
	non_ministry = [n for n in frappe.get_all(DOCTYPE, pluck="name", limit_page_length=0) if n not in ministry]
	if not non_ministry:
		print("Nothing to do: every code is a Ministry code.")
		return

	usages, unmapped = [], set()
	for doctype, fieldname in lookup_fields():
		if not frappe.db.table_exists(doctype) or not frappe.db.has_column(doctype, fieldname):
			continue
		for value, count in frappe.db.sql(
			f"select `{fieldname}`, count(*) from `tab{doctype}` where `{fieldname}` in %s group by 1",
			(tuple(non_ministry),),
		):
			replacement = DEMO_REPLACEMENTS.get(value)
			if not replacement or replacement not in ministry:
				unmapped.add(value)
			usages.append((doctype, fieldname, value, replacement, count))

	print(f"{len(non_ministry)} codes are not in the Ministry dataset.")
	for doctype, fieldname, value, replacement, count in usages:
		print(f"  {doctype}.{fieldname}: {value} -> {replacement or 'NO REPLACEMENT'} ({count})")
	if unmapped:
		print(f"Stopped: add these to DEMO_REPLACEMENTS first: {sorted(unmapped)}")
		return
	if not commit:
		print("DRY RUN: nothing changed. Run with commit=1 to apply.")
		return

	for doctype, fieldname, value, replacement, _count in usages:
		frappe.db.sql(f"update `tab{doctype}` set `{fieldname}`=%s where `{fieldname}`=%s", (replacement, value))
	for name in non_ministry:
		frappe.delete_doc(DOCTYPE, name, ignore_permissions=True, force=True)
	frappe.db.commit()
	frappe.clear_cache()
	print(f"Replaced {sum(u[4] for u in usages)} references and deleted {len(non_ministry)} non-Ministry codes.")
