import re
from pathlib import Path

import frappe
from openpyxl import load_workbook


DOCTYPE = "Jameah Ministry Code"


def _normalize_category(sheet_name: str) -> str:
	name = re.sub(r"^\s*\d+\s*[-.)]\s*", "", sheet_name or "")
	name = re.sub(r"\s+", " ", name).strip()
	return name or "General"


def _text(value) -> str:
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	return str(value).strip()


def _norm_header_token(value) -> str:
	text = _text(value).lower()
	text = re.sub(r"[-_]+", " ", text)
	text = re.sub(r"[^a-z0-9\u0600-\u06ff\s]+", "", text)
	return re.sub(r"\s+", " ", text).strip()


def _is_header_like(value: str) -> bool:
	v = _norm_header_token(value)
	if not v:
		return True
	blocked = {
		"symbol",
		"the symbol",
		"name",
		"the name",
		"english name",
		"name in english",
		"coding",
		"last modified date",
		"edited by",
		"الرمز",
		"الاسم",
	}
	return v in blocked


def _is_code_like(value: str) -> bool:
	v = re.sub(r"\s+", "", _text(value).upper())
	if not v:
		return False
	if len(v) > 15:
		return False
	if not re.search(r"\d", v):
		return False
	return bool(re.fullmatch(r"[A-Z0-9][A-Z0-9\-/]{0,14}", v))


def _find_header_map(sheet):
	max_scan_rows = min(12, sheet.max_row)
	best = None

	for row_idx in range(1, max_scan_rows + 1):
		row = [sheet.cell(row=row_idx, column=col).value for col in range(1, sheet.max_column + 1)]
		tokens = [_norm_header_token(v) for v in row]

		symbol_col = None
		coding_col = None
		name_ar_col = None
		name_en_col = None

		for col_idx, token in enumerate(tokens):
			if not token:
				continue
			if symbol_col is None and ("symbol" in token or "الرمز" in token):
				symbol_col = col_idx
			if coding_col is None and "coding" in token:
				coding_col = col_idx
			if name_en_col is None and ("name in english" in token or "english name" in token):
				name_en_col = col_idx
			if name_ar_col is None and (token == "name" or token == "the name" or "name in arabic" in token or "الاسم" in token):
				name_ar_col = col_idx

		score = sum(v is not None for v in [symbol_col, coding_col, name_ar_col, name_en_col])
		if score == 0:
			continue

		candidate = {
			"header_row": row_idx,
			"symbol_col": symbol_col,
			"coding_col": coding_col,
			"name_ar_col": name_ar_col,
			"name_en_col": name_en_col,
			"score": score,
		}

		if not best or candidate["score"] > best["score"]:
			best = candidate

	if best:
		return best

	return {
		"header_row": 1,
		"symbol_col": 0,
		"coding_col": None,
		"name_ar_col": 1,
		"name_en_col": 2,
		"score": 0,
	}


def _pick_code(symbol_value, coding_value):
	symbol = _text(symbol_value)
	coding = _text(coding_value)
	if symbol and not _is_header_like(symbol) and _is_code_like(symbol):
		return symbol.upper()
	if coding and not _is_header_like(coding) and _is_code_like(coding):
		return coding.upper()
	return ""


def execute(file_path: str, dry_run: bool = False):
	path = Path(file_path)
	if not path.exists():
		frappe.throw(f"Excel file not found: {file_path}")

	wb = load_workbook(filename=str(path), data_only=True)
	updated = 0
	skipped = 0
	not_found = 0

	for sheet_name in wb.sheetnames:
		sheet = wb[sheet_name]
		category = _normalize_category(sheet_name)
		header = _find_header_map(sheet)

		for row in sheet.iter_rows(min_row=header["header_row"] + 1, values_only=True):
			symbol = (row[header["symbol_col"]] if header["symbol_col"] is not None and len(row) > header["symbol_col"] else None)
			coding = (row[header["coding_col"]] if header["coding_col"] is not None and len(row) > header["coding_col"] else None)
			code = _pick_code(symbol, coding)
			if not code:
				skipped += 1
				continue

			name_ar_raw = (row[header["name_ar_col"]] if header["name_ar_col"] is not None and len(row) > header["name_ar_col"] else None)
			name_en_raw = (row[header["name_en_col"]] if header["name_en_col"] is not None and len(row) > header["name_en_col"] else None)

			name_en = _text(name_en_raw)
			name_ar = _text(name_ar_raw)

			if _is_header_like(name_en):
				name_en = ""
			if _is_header_like(name_ar):
				name_ar = ""

			# In many sheets, the visible name is placed in the "name" column.
			if not name_en and name_ar:
				name_en = name_ar

			if not name_en:
				name_en = code

			# Placeholder Arabic fallback until machine translation is added.
			if not name_ar:
				name_ar = name_en

			docname = frappe.db.get_value(DOCTYPE, {"code_category": category, "ministry_code": code}, "name")
			if not docname:
				not_found += 1
				continue

			if dry_run:
				updated += 1
				continue

			frappe.db.set_value(
				DOCTYPE,
				docname,
				{"name_english": name_en, "name_arabic": name_ar},
				update_modified=False,
			)
			updated += 1

	if not dry_run:
		frappe.db.commit()
		frappe.clear_cache(doctype=DOCTYPE)

	print(f"Fix names from file: {file_path}")
	print(f"Mode: {'dry_run' if dry_run else 'write'}")
	print(f"Updated: {updated}")
	print(f"Skipped: {skipped}")
	print(f"Not found by category+code: {not_found}")
