"""HR spreadsheet to complete the Ministry (Jameah) data of non-academic employees.

The employees listed are the ones the Ministry Employee views send (teaching staff go through
the Instructor data sheet and are sent as Faculty). Current values are filled in; empty
Ministry-mandatory cells are marked red. HR fills the sheet, then checks it (nothing is saved,
a result sheet gives each row's status) and imports it. Rows are saved through the normal
Employee validations, so a row the Employee form would refuse is reported, not saved.

From the "Employees Missing Ministry Data" report (HR Manager): Download / Check / Import, or:
    bench --site <site> execute medical_hrms.employee_data_sheet.export_sheet
    bench --site <site> execute medical_hrms.employee_data_sheet.import_sheet --kwargs "{'file': '/private/files/<name>.xlsx'}"
    ... import_sheet --kwargs "{'file': '/private/files/<name>.xlsx', 'commit': 1}"
"""
import re

import frappe
from frappe import _
from frappe.utils import cstr, getdate, validate_email_address

from medical_hrms.instructor_data_sheet import INPUT_FILL, LOCKED_FILL, _resolve_code, _save_file
from medical_hrms.jameah.views import EMPLOYEE_POPULATION
from medical_hrms.ministry_lookup_validation import ministry_codes

FILE_PREFIX = "Employee-Ministry-Data"
SHEET = "Employees"
MISSING_FILL = "F8CBAD"  # red: Ministry-mandatory and empty
ROLE = "HR Manager"

# (key, header, width, kind, Employee field, Ministry code category or None, mandatory for the Ministry)
COLUMNS = [
	("employee", "Employee ID / رقم الموظف [employee]", 16, "locked", None, None, False),
	("employee_name", "Name / الاسم [employee_name]", 28, "locked", None, None, False),
	("department", "Department / القسم [department]", 24, "locked", None, None, False),
	("first_name_ar", "First Name (Arabic) / الاسم الأول [first_name_ar]", 18, "input", "custom_first_name_ar", None, True),
	("second_name_ar", "Second Name (Arabic) / اسم الأب [second_name_ar]", 18, "input", "custom_second_name_ar", None, True),
	("third_name_ar", "Third Name (Arabic) / اسم الجد [third_name_ar]", 18, "input", "custom_third_name_ar", None, False),
	("last_name_ar", "Last Name (Arabic) / اسم العائلة [last_name_ar]", 18, "input", "custom_last_name_ar", None, True),
	("first_name_en", "First Name (English) [first_name_en]", 18, "input", "custom_first_name_en", None, True),
	("second_name_en", "Second Name (English) [second_name_en]", 18, "input", "custom_second_name_en", None, True),
	("third_name_en", "Third Name (English) [third_name_en]", 18, "input", "custom_third_name_en", None, False),
	("last_name_en", "Last Name (English) [last_name_en]", 18, "input", "custom_last_name_en", None, True),
	("gender", "Gender / الجنس [gender]", 12, "input", "gender", "Gender", True),
	("date_of_birth", "Date of Birth (YYYY-MM-DD) / تاريخ الميلاد [date_of_birth]", 16, "input", "date_of_birth", None, True),
	("nationality", "Nationality / الجنسية [nationality]", 20, "input", "custom_jameah_nationality", "Nationality", True),
	("identity_type", "Identity Type / نوع الهوية [identity_type]", 16, "input", "custom_identity_type", "Identity type", True),
	("identity_number", "Identity Number (National ID / Iqama: 10 digits) / رقم الهوية [identity_number]", 22, "input", "custom_identity_number", None, True),
	("home_id", "Optional: ID in Home Country (non-Saudis) / رقم الهوية في بلد الأصل [home_id]", 22, "input", "custom_original_home_id_number", None, False),
	("work_city", "City of Work / مدينة العمل [work_city]", 22, "input", "custom_ministry_work_city", "Coding cities and governora", True),
	("mobile", "Mobile / الجوال [mobile]", 16, "input", "cell_number", None, True),
	("email", "Email / البريد [email]", 28, "input", "prefered_email", None, True),
	("special_needs", "Special Needs (Yes / No) / ذوي الاحتياجات الخاصة [special_needs]", 14, "input", "custom_is_special_needs", None, True),
	("special_needs_type", "Type of Special Needs (if Yes) / نوع الاحتياج [special_needs_type]", 22, "input", "custom_ministry_special_needs_type", "Type of special needs", False),
]
KEYS = [c[0] for c in COLUMNS]
# What the Ministry views fall back to when the field is empty (jameah/views.py), so the sheet shows
# - and only marks as missing - what the Ministry actually receives.
FALLBACK = {
	"custom_first_name_en": ("first_name",), "custom_second_name_en": ("middle_name",), "custom_last_name_en": ("last_name",),
	"prefered_email": ("company_email", "personal_email"),
}


def _current(employee, field):
	for f in (field, *FALLBACK.get(field, ())):
		value = cstr(employee.get(f)).strip()
		if value:
			return value
	return ""
TEXT_COLUMNS = {"identity_number", "home_id", "mobile"}
YES, NO = "Yes", "No"

INSTRUCTIONS = [
	"Employee Ministry Data / بيانات الموظفين للوزارة",
	"",
	"Why: the Ministry (Jameah) views send these values for every non-academic employee. Empty ones are sent empty.",
	"Teaching staff are not listed here: they are sent as Faculty (use the Instructor data sheet for them).",
	"",
	"How to fill in the 'Employees' sheet:",
	"  • Grey columns: do not change.",
	"  • Green columns: correct or fill in. Red cells are required by the Ministry and still empty.",
	"  • Leave a cell empty to keep what the Employee record has (an empty cell never deletes data).",
	"  • Identity: National ID for Saudis (10 digits starting with 1), Iqama for residents (10 digits starting with 2).",
	"  • Use the dropdowns where available. Dates as YYYY-MM-DD.",
	"",
	"Then: Employees Missing Ministry Data report > Check Data Sheet (nothing is saved, you get a result sheet),",
	"fix what it reports, and Import Data Sheet.",
]


def _require_hr_manager():
	if frappe.session.user != "Administrator" and ROLE not in frappe.get_roles():
		frappe.throw(_("Only an HR Manager can export or import the Employee data sheet."), frappe.PermissionError)


def _dated_name(step):
	return f"{frappe.utils.now_datetime():%Y-%m-%d_%H-%M}_{FILE_PREFIX}_{step}.xlsx"


def _employees():
	return frappe.db.sql(f"SELECT e.* FROM {EMPLOYEE_POPULATION} ORDER BY e.employee_name", as_dict=True)


def _codes():
	return {cat: ministry_codes(cat) for *_rest, cat, _m in COLUMNS if cat}


@frappe.whitelist()
def export_sheet():
	_require_hr_manager()
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.worksheet.datavalidation import DataValidation

	codes = _codes()
	label = {c.name: c.name_english for options in codes.values() for c in options}
	employees = _employees()

	workbook = Workbook()
	guide = workbook.active
	guide.title = "Instructions"
	for row, line in enumerate(INSTRUCTIONS, start=1):
		guide.cell(row=row, column=1, value=line)
	guide.column_dimensions["A"].width = 115

	sheet = workbook.create_sheet(SHEET)
	fills = {"locked": PatternFill("solid", fgColor=LOCKED_FILL), "input": PatternFill("solid", fgColor=INPUT_FILL)}
	missing_fill = PatternFill("solid", fgColor=MISSING_FILL)
	for col, (key, header, width, kind, *_rest) in enumerate(COLUMNS, start=1):
		cell = sheet.cell(row=1, column=col, value=header)
		cell.font = Font(bold=True)
		cell.alignment = Alignment(wrap_text=True, vertical="center")
		cell.fill = fills[kind]
		sheet.column_dimensions[cell.column_letter].width = width
	sheet.row_dimensions[1].height = 52
	sheet.freeze_panes = "D2"

	missing_cells = 0
	for row, e in enumerate(employees, start=2):
		for col, (key, _h, _w, kind, field, category, mandatory) in enumerate(COLUMNS, start=1):
			if key in ("employee", "employee_name", "department"):
				value = e.get("name" if key == "employee" else key)
			elif key == "special_needs":
				value = YES if e.get(field) else NO
			elif category:
				value = label.get(e.get(field)) or None
			elif key == "date_of_birth":
				value = cstr(e.get(field)) or None
			else:
				value = _current(e, field) or None
			cell = sheet.cell(row=row, column=col, value=value)
			if key in TEXT_COLUMNS or key == "date_of_birth":
				cell.number_format = "@"
			if mandatory and not value:
				cell.fill = missing_fill
				missing_cells += 1

	lists = workbook.create_sheet("Lists")
	lists.sheet_state = "hidden"
	sources = {key: list(dict.fromkeys(c.name_english for c in codes[category] if c.name_english))
		for key, *_m, category, _x in COLUMNS if category}
	sources["special_needs"] = [YES, NO]
	last = max(len(employees) + 1, 2)
	for list_col, (key, options) in enumerate(sources.items(), start=1):
		for i, option in enumerate(options, start=1):
			lists.cell(row=i, column=list_col, value=option)
		letter = lists.cell(row=1, column=list_col).column_letter
		target = sheet.cell(row=2, column=KEYS.index(key) + 1).column_letter
		validation = DataValidation(type="list", formula1=f"=Lists!${letter}$1:${letter}${max(len(options), 1)}")
		sheet.add_data_validation(validation)
		validation.add(f"{target}2:{target}{last}")

	file_url = _save_file(_dated_name("1-For-HR"), workbook)
	return {"file_url": file_url, "employees": len(employees), "missing_cells": missing_cells}


def _read_rows(file):
	from openpyxl import load_workbook

	is_file_url = file.startswith(("/private/files/", "/files/"))
	path = frappe.get_doc("File", {"file_url": file}).get_full_path() if is_file_url else file
	sheet = load_workbook(path, data_only=True)[SHEET]
	header_keys = []
	for cell in sheet[1]:
		match = re.search(r"\[(\w+)\]\s*$", cstr(cell.value))
		header_keys.append(match.group(1) if match else None)
	for excel_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
		values = {}
		for key, value in zip(header_keys, row):
			if not key:
				continue
			if isinstance(value, float) and value.is_integer():
				value = int(value)  # Excel turns typed IDs into numbers
			if hasattr(value, "strftime"):
				value = value.strftime("%Y-%m-%d")
			values[key] = cstr(value).strip() if value is not None else ""
		if values.get("employee"):
			yield excel_row, values


def _changes(values, employee, codes):
	"""(changes for the Employee, problems). An empty cell keeps the current value."""
	changes, problems = {}, []
	for key, _h, _w, kind, field, category, _m in COLUMNS:
		value = values.get(key, "")
		if kind == "locked" or not value:
			continue
		if category:
			code = _resolve_code(value, codes[category])
			if not code:
				problems.append(_("{0} '{1}' is not in the Ministry list").format(key.replace("_", " ").title(), value))
				continue
			value = code
		elif key == "date_of_birth":
			try:
				value = str(getdate(value))
			except Exception:
				problems.append(_("Date of Birth '{0}' is not a date (use YYYY-MM-DD)").format(value))
				continue
		elif key == "special_needs":
			if value not in (YES, NO):
				problems.append(_("Special Needs must be Yes or No"))
				continue
			value = 1 if value == YES else 0
		elif key == "mobile":
			if not re.fullmatch(r"\+?[0-9 ]{9,15}", value):
				problems.append(_("Mobile '{0}' is not a phone number").format(value))
				continue
		elif key == "email":
			if not validate_email_address(value):
				problems.append(_("Email '{0}' is not valid").format(value))
				continue
		current = _current(employee, field) if field in FALLBACK else cstr(employee.get(field))
		if current != cstr(value):
			changes[field] = value
	return changes, problems


@frappe.whitelist()
def import_sheet(file, commit=0):
	_require_hr_manager()
	from openpyxl import Workbook
	from openpyxl.styles import PatternFill

	commit = int(commit)
	codes = _codes()
	allowed = {e.name for e in _employees()}
	results = []
	for excel_row, values in _read_rows(file):
		name = values["employee"]
		if name not in allowed:
			results.append((excel_row, values, "Fix needed", _("Employee {0} is not one of the employees sent to the Ministry (grey columns changed?)").format(name)))
			continue
		employee = frappe.get_doc("Employee", name)
		changes, problems = _changes(values, employee, codes)
		if problems:
			results.append((excel_row, values, "Fix needed", "; ".join(problems)))
			continue
		if not changes:
			results.append((excel_row, values, "No change", ""))
			continue
		frappe.db.savepoint("employee_sheet_row")
		try:
			employee.update(changes)
			employee.save()
		except Exception as e:
			frappe.db.rollback(save_point="employee_sheet_row")
			frappe.clear_messages()
			message = re.sub(r"<[^>]+>", " ", cstr(e)).strip() or e.__class__.__name__
			results.append((excel_row, values, "Fix needed", message))
			continue
		label = ", ".join(employee.meta.get_label(f) for f in changes)
		results.append((excel_row, values, "Done" if commit else "Ready", (_("Updated: {0}") if commit else _("Will update: {0}")).format(label)))

	if not commit:
		frappe.db.rollback()

	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "Result"
	sheet.append(["Excel Row", "Employee ID", "Name", "Status / الحالة", "Details / التفاصيل"])
	colours = {"Ready": "C6EFCE", "Done": "C6EFCE", "Fix needed": "FFC7CE", "No change": "EDEDED"}
	for excel_row, values, status, details in results:
		sheet.append([excel_row, values["employee"], values.get("employee_name"), status, details])
		sheet.cell(row=sheet.max_row, column=4).fill = PatternFill("solid", fgColor=colours[status])
	for letter, width in zip("ABCDE", (10, 16, 30, 16, 100)):
		sheet.column_dimensions[letter].width = width
	sheet.freeze_panes = "A2"
	file_url = _save_file(_dated_name("3-Import-Done" if commit else "2-Check-Result"), workbook)

	counts = {}
	for *_r, status, _d in results:
		counts[status] = counts.get(status, 0) + 1
	return {"result_file": file_url, "counts": counts, "saved": bool(commit)}
