import re
from pathlib import Path

import frappe
from openpyxl import load_workbook


DOCTYPE = "Jameah Ministry Code"


def _normalize_category(sheet_name: str) -> str:
	name = re.sub(r"^\s*\d+\s*[-.)]\s*", "", sheet_name or "")
	name = re.sub(r"\s+", " ", name).strip()
	return name or "General"


def _text(value) -> str:
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	return str(value).strip()


def _norm_header_token(value) -> str:
	text = _text(value).lower()
	text = re.sub(r"[-_]+", " ", text)
	text = re.sub(r"[^a-z0-9\u0600-\u06ff\s]+", "", text)
	return re.sub(r"\s+", " ", text).strip()


def _find_header_map(sheet):
	max_scan_rows = min(12, sheet.max_row)
	best = None

	for row_idx in range(1, max_scan_rows + 1):
		row = [sheet.cell(row=row_idx, column=col).value for col in range(1, sheet.max_column + 1)]
		tokens = [_norm_header_token(v) for v in row]

		symbol_col = None
		coding_col = None
		name_ar_col = None
		name_en_col = None

		for col_idx, token in enumerate(tokens):
			if not token:
				continue
			if symbol_col is None and ("symbol" in token or "الرمز" in token):
				symbol_col = col_idx
			if coding_col is None and "coding" in token:
				coding_col = col_idx
			if name_en_col is None and "name in english" in token:
				name_en_col = col_idx
			if name_ar_col is None and (token == "name" or "name in arabic" in token or "الاسم" in token):
				name_ar_col = col_idx

		score = sum(v is not None for v in [symbol_col, coding_col, name_ar_col, name_en_col])
		if score == 0:
			continue

		candidate = {
			"header_row": row_idx,
			"symbol_col": symbol_col,
			"coding_col": coding_col,
			"name_ar_col": name_ar_col,
			"name_en_col": name_en_col,
			"score": score,
		}

		if not best or candidate["score"] > best["score"]:
			best = candidate

	if best:
		return best

	# Fallback to legacy layout: A=symbol, B=arabic, C=english
	return {
		"header_row": 1,
		"symbol_col": 0,
		"coding_col": None,
		"name_ar_col": 1,
		"name_en_col": 2,
		"score": 0,
	}


def _is_header_like(value: str) -> bool:
	if not value:
		return True
	v = _norm_header_token(value)
	if not v:
		return True

	blocked_tokens = [
		"symbol",
		"the symbol",
		"name",
		"name in english",
		"coding",
		"last modified date",
		"edited by",
		"الرمز",
		"الاسم",
	]

	return v in blocked_tokens


def _is_code_like(value: str) -> bool:
	if not value:
		return False
	v = re.sub(r"\s+", "", _text(value).upper())
	if not v or len(v) > 15:
		return False
	if not re.search(r"\d", v):
		return False
	# Allow alphanumeric ministry codes and optional -/
	return bool(re.fullmatch(r"[A-Z0-9][A-Z0-9\-/]{0,14}", v))


def _pick_code(symbol_value, coding_value):
	symbol = _text(symbol_value)
	coding = _text(coding_value)

	if symbol and not _is_header_like(symbol) and _is_code_like(symbol):
		return symbol.upper(), "symbol"

	if coding and not _is_header_like(coding) and _is_code_like(coding):
		return coding.upper(), "coding"

	return "", "none"


def _looks_like_section_title(value: str) -> bool:
	v = _text(value)
	if not v:
		return False
	if _is_code_like(v):
		return False
	if len(v) < 8:
		return False
	if len(v.split()) >= 3:
		return True
	return False


def _upsert_code(category: str, ministry_code: str, name_english: str, name_arabic: str) -> str:
	existing = frappe.db.get_value(
		DOCTYPE,
		{"code_category": category, "ministry_code": ministry_code},
		"name",
	)

	values = {
		"code_category": category,
		"ministry_code": ministry_code,
		"name_english": name_english or ministry_code,
		"name_arabic": name_arabic or name_english or ministry_code,
	}

	if existing:
		frappe.db.set_value(DOCTYPE, existing, values, update_modified=False)
		return "updated"

	doc = frappe.get_doc({"doctype": DOCTYPE, **values})
	doc.insert(ignore_permissions=True)
	return "inserted"


def execute(file_path: str, sheet_names: list[str] | None = None, dry_run: bool = False):
	path = Path(file_path)
	if not path.exists():
		frappe.throw(f"Excel file not found: {file_path}")

	wb = load_workbook(filename=str(path), data_only=True)
	target_sheets = sheet_names or wb.sheetnames

	inserted = 0
	updated = 0
	skipped = 0
	rejected = 0
	used_symbol = 0
	used_coding = 0
	sheet_stats = []

	for sheet_name in target_sheets:
		if sheet_name not in wb.sheetnames:
			print(f"Skipped missing sheet: {sheet_name}")
			continue

		sheet = wb[sheet_name]
		category = _normalize_category(sheet_name)
		header = _find_header_map(sheet)

		sheet_inserted = 0
		sheet_updated = 0
		sheet_skipped = 0
		sheet_rejected = 0

		for row in sheet.iter_rows(min_row=header["header_row"] + 1, values_only=True):
			symbol = (row[header["symbol_col"]] if header["symbol_col"] is not None and len(row) > header["symbol_col"] else None)
			coding = (row[header["coding_col"]] if header["coding_col"] is not None and len(row) > header["coding_col"] else None)
			name_ar = (row[header["name_ar_col"]] if header["name_ar_col"] is not None and len(row) > header["name_ar_col"] else None)
			name_en = (row[header["name_en_col"]] if header["name_en_col"] is not None and len(row) > header["name_en_col"] else None)

			ministry_code, source = _pick_code(symbol, coding)
			if not ministry_code:
				sheet_skipped += 1
				skipped += 1
				continue

			name_english = _text(name_en)
			name_arabic = _text(name_ar)

			# Guard against section/title rows being imported as records.
			if _looks_like_section_title(name_english) and ministry_code == name_english.upper().replace(" ", ""):
				sheet_rejected += 1
				rejected += 1
				continue
			if _looks_like_section_title(name_arabic) and ministry_code == name_arabic.upper().replace(" ", ""):
				sheet_rejected += 1
				rejected += 1
				continue

			if source == "symbol":
				used_symbol += 1
			elif source == "coding":
				used_coding += 1

			if dry_run:
				sheet_inserted += 1
				inserted += 1
				continue

			result = _upsert_code(category, ministry_code, name_english, name_arabic)
			if result == "inserted":
				sheet_inserted += 1
				inserted += 1
			else:
				sheet_updated += 1
				updated += 1

		sheet_stats.append((sheet_name, sheet_inserted, sheet_updated, sheet_skipped, sheet_rejected))

	if not dry_run:
		frappe.db.commit()
		frappe.clear_cache(doctype=DOCTYPE)

	print(f"Imported from file: {file_path}")
	print(f"Mode: {'dry_run' if dry_run else 'write'}")
	print(f"Inserted: {inserted}")
	print(f"Updated: {updated}")
	print(f"Skipped: {skipped}")
	print(f"Rejected: {rejected}")
	print(f"Code source usage -> symbol: {used_symbol}, coding: {used_coding}")
	for sheet_name, s_i, s_u, s_s, s_r in sheet_stats:
		print(f"Sheet {sheet_name}: inserted={s_i}, updated={s_u}, skipped={s_s}, rejected={s_r}")
